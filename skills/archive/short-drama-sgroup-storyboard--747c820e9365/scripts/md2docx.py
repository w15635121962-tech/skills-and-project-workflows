#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把分镜 Markdown 转成可打开的 Word 文档（python-docx）。
支持：标题 / 段落 / 引用 / 表格 / 代码块 / 水平线 / 加粗 / 行内代码 / 无序列表
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK = "微软雅黑"
MONO = "Consolas"

DARK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x8B, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
CODE_FG = RGBColor(0x1F, 0x33, 0x22)


def set_cjk(run, font=CJK):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), font)
    rf.set(qn("w:ascii"), font)
    rf.set(qn("w:hAnsi"), font)


def shade(el, hexcolor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    el.append(shd)


def para_border(p, position="bottom", color="CCCCCC", size=6):
    ppr = p._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    b = OxmlElement("w:" + position)
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), color)
    pbdr.append(b)


def add_rich(p, text, base_size=10.5, base_color=DARK, bold_all=False):
    """解析 **bold** 与 `code` 行内标记。"""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("**") and tk.endswith("**") and len(tk) > 4:
            r = p.add_run(tk[2:-2])
            r.bold = True
        elif tk.startswith("`") and tk.endswith("`") and len(tk) > 2:
            r = p.add_run(tk[1:-1])
            r.font.name = MONO
            r._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
            rf = r._element.rPr.find(qn("w:rFonts"))
            rf.set(qn("w:eastAsia"), MONO)
            rf.set(qn("w:ascii"), MONO)
            rf.set(qn("w:hAnsi"), MONO)
            r.font.size = Pt(base_size - 1)
            r.font.color.rgb = RGBColor(0x0B, 0x4F, 0x6C)
            continue
        else:
            r = p.add_run(tk)
        if bold_all:
            r.bold = True
        set_cjk(r)
        r.font.size = Pt(base_size)
        r.font.color.rgb = base_color


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def build(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    st = doc.styles["Normal"]
    st.font.name = CJK
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.18

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i]
        line = raw.rstrip()

        # 代码块
        if line.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            shade(p._p.get_or_add_pPr(), "F4F6F4")
            para_border(p, "left", "9AA79A", 18)
            r = p.add_run("\n".join(buf))
            r.font.name = MONO
            rf = r._element.get_or_add_rPr().find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                r._element.rPr.append(rf)
            rf.set(qn("w:eastAsia"), MONO)
            rf.set(qn("w:ascii"), MONO)
            rf.set(qn("w:hAnsi"), MONO)
            r.font.size = Pt(9)
            r.font.color.rgb = CODE_FG
            continue

        # 表格
        if line.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = split_row(line)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            ncol = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
            t = doc.add_table(rows=1, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = t.rows[0].cells
            for c in range(ncol):
                txt = header[c] if c < len(header) else ""
                cp = hdr[c].paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_rich(cp, txt, 9.5, DARK, bold_all=True)
                shade(hdr[c]._tc.get_or_add_tcPr(), "EDEFED")
            for rw in rows:
                cells = t.add_row().cells
                for c in range(ncol):
                    txt = rw[c] if c < len(rw) else ""
                    cp = cells[c].paragraphs[0]
                    cp.paragraph_format.space_after = Pt(0)
                    add_rich(cp, txt, 9.5, DARK)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 水平线
        if re.match(r"^-{3,}$", line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            para_border(p, "bottom", "BBBBBB", 6)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            size = {1: 19, 2: 15.5, 3: 13, 4: 11.5}.get(lvl, 11)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            add_rich(p, txt, size, ACCENT if lvl <= 3 else DARK, bold_all=True)
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            txt = line.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            para_border(p, "left", "C9A227", 14)
            add_rich(p, txt, 9.5, MUTED)
            i += 1
            continue

        # 无序列表
        if re.match(r"^\s*[-*]\s+", line):
            txt = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, txt, 10.5, DARK)
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_rich(p, line.strip(), 10.5, DARK)
        i += 1
        continue

    doc.save(docx_path)
    print("已生成:", docx_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
