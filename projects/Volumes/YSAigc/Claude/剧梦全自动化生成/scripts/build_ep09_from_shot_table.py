#!/usr/bin/env python3
"""Build the EP09 storyboard/prompt artifacts from the approved shot table."""
from __future__ import annotations
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import automator

TABLE = ROOT / "01自动化项目素材/玻璃箱的神明/玻璃箱里的神明分镜拆解/玻璃箱里的神明-全剧镜头表.xlsx"
RUN = ROOT / "runs/20260918-181445-a2e312"

def quote_only(text: str) -> str:
    values = re.findall(r'[“"]([^”"]+)[”"]', text or "")
    return " ".join(v.strip() for v in values if v.strip())

def shot_scene(n: int) -> str:
    if n <= 3: return "商场闪回与恒温箱"
    if n <= 12: return "Blackwood庄园客厅"
    return "主卧"

def characters(n: int):
    if n <= 3: return ["Hades 小黑蛇"]
    if n in (6, 7, 12, 17, 18, 19, 24, 26, 28, 29): return ["Aria"]
    if n in (14, 15, 16, 20, 21, 22, 23, 25, 27, 30, 31): return ["Aria", "Hades"]
    return ["Hades"]

def refs(chars):
    out=[]
    if "Aria" in chars: out.append("ep09_character_aria")
    if "Hades" in chars: out.append("ep09_character_hades_human")
    if "Hades 小黑蛇" in chars: out.append("ep09_character_hades_snake")
    return out

def main():
    ws = load_workbook(TABLE, read_only=True, data_only=True)["全剧镜头表"]
    shots=[]
    for row in ws.iter_rows(values_only=True):
        if str(row[0]).strip() != "EP09": continue
        n=int(row[3]); duration=float(row[4]); camera=str(row[5] or ""); action=str(row[6] or ""); english=str(row[7] or ""); performance=str(row[8] or "")
        chars=characters(n); dialogue=quote_only(english)
        comp=automator.composition_parameters(camera)
        shots.append({
            "id": f"EP09-S{n:02d}", "duration": duration, "task": "EP09 approved shot-table beat",
            "characters": chars, "scene": shot_scene(n), "action": action,
            "camera": camera, "composition": action, "blocking": "严格按批准镜头表完成动作，不新增角色、道具或结果。",
            "eyeline": "视线只落在明确的戏内目标；正反打保持本镜既定轴线与屏幕方向。",
            "performance": performance, "dialogue": dialogue, "english_dialogue": dialogue,
            "asset_refs": refs(chars), "doubao_design": {"source":"approved_shot_table", "negative":"保持镜头表的角色、动作、道具和终态，不添加现实标识。"},
            "lighting": "延续本场既定光线与色温；转场只改变光影，不改变角色身份与背景锚点。",
            "audio": "现场环境音、动作音效；仅在本镜有英文对白时保留原文，不生成旁白或字幕。",
            "motion": "按批准镜头表执行，切点落在动作进行中。", **comp,
            **automator.narrative_story_design(action)
        })
    sb={"model":"豆包2.1Turbo", "timing_source":"approved_EP09_shot_table", "axis":"EP09连续镜头保持批准镜头表轴线与屏幕方向", "shots":shots}
    report=automator.qa({}, {"missing":[]}, sb, {"quality":{"require_asset_binding":True}}, None)
    (RUN/"storyboard.json").write_text(json.dumps(sb, ensure_ascii=False, indent=2), encoding="utf-8")
    (RUN/"qa_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (RUN/"format_c_prompts.md").write_text(automator.format_c(sb, "剧梦2.5Pro", "480P"), encoding="utf-8")
    print(json.dumps({"shots":len(shots),"duration":sum(s["duration"] for s in shots),"qa":report["status"],"run":str(RUN)}, ensure_ascii=False))

if __name__ == "__main__": main()
